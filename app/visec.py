"""Mapeo de una CPE de ARCA a la fila del XLSX de Visec, y escritura del archivo.

La plantilla (`plantilla_visec.xlsx`) es la que entrega Visec y no se modifica:
solo se le agregan filas. Las columnas se verifican antes de escribir, así que
si Visec cambia la plantilla el proceso se detiene en vez de generar un archivo
que el portal de ellos va a rechazar.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import PLANTILLA, config
from .wscpe import Consulta

COLUMNAS = [
    "Fecha y Hora Movimiento", "Fecha CPE", "Número CPE", "Número CTG",
    "CUIT Titular", "Número RUCA Origen", "CUIT Remitente Comercial Productor",
    "CUIT Rte Comercial Venta Primaria", "CUIT Rte Comercial Venta Secundaria",
    "CUIT Rte Comercial Venta Secundaria 2", "CUIT Corredor Venta Primaria",
    "CUIT Corredor Venta Secundaria", "CUIT Destinatario", "CUIT Destino",
    "Número RUCA Destino", "Código Producto", "Campaña", "Peso Neto Carga (Kg)",
    "Número RENSPA", "Número CTG Asignado", "Peso Neto Carga (Kg) por UP",
    "Peso Neto Descarga (Kg) por UP", "Peso Ingreso Stock (Kg)",
    "Último Almacenamiento", "Tipo Movimiento",
]

HOJA = "Cartas de Porte"

# Estados de la CPE (manual wscpe v2.0.5, 2.7.1).
ESTADOS = {
    "AC": "Activa",
    "AN": "Anulada",
    "AP": "Anulada por el productor",
    "BR": "Borrador",
    "CF": "Activa con confirmación de arribo",
    "CN": "Confirmada",
    "CO": "Activa con contingencia",
    "DD": "Descargada en destino",
    "DE": "Desactivada",
    "IN": "Inactiva",
    "PA": "Pendiente de aceptación del productor",
    "PE": "Pendiente de emisión",
    "PO": "Pendiente de aceptación del origen",
    "RE": "Rechazada",
}

# Solo una CPE cerrada en destino tiene los pesos de descarga definitivos. En
# cualquier otro estado puede cambiar o no haber llegado todavía.
ESTADOS_FINALES = {"CN", "DD"}


# ---------------------------------------------------------------------------
# Lectura de la respuesta de ARCA
# ---------------------------------------------------------------------------

def _buscar(datos: Any, ruta: str) -> Any:
    """Valor en una ruta con puntos: _buscar(cpe, 'destino.planta')."""
    actual = datos
    for parte in ruta.split("."):
        if isinstance(actual, list):
            actual = actual[0] if actual else None
        if not isinstance(actual, dict):
            return None
        actual = actual.get(parte)
    return actual[0] if isinstance(actual, list) and actual else actual


def _texto(datos: Any, ruta: str) -> str:
    valor = _buscar(datos, ruta)
    return "" if valor is None else str(valor).strip()


def _cuit(datos: Any, ruta: str) -> str:
    """CUIT en dígitos; '' si no tiene 11 (ARCA devuelve 0 cuando no aplica)."""
    digitos = "".join(c for c in _texto(datos, ruta) if c.isdigit())
    return digitos if len(digitos) == 11 else ""


def _numero(valor: Any) -> Decimal | None:
    if valor is None or valor == "":
        return None
    try:
        return Decimal(str(valor).replace(",", "."))
    except InvalidOperation:
        return None


def _kilos(datos: Any, ruta_bruto: str, ruta_tara: str) -> str:
    """Peso neto = bruto - tara, redondeado a kilos enteros."""
    bruto, tara = _numero(_buscar(datos, ruta_bruto)), _numero(_buscar(datos, ruta_tara))
    if bruto is None:
        return ""
    neto = bruto - (tara or 0)
    return str(int(neto.to_integral_value())) if neto > 0 else ""


def _kilos_propios(valor: str | None) -> str:
    """Kilos que vienen del CSV del sistema, redondeados a entero."""
    numero = _numero(valor)
    return str(int(numero.to_integral_value())) if numero and numero > 0 else ""


def _fecha(valor: Any, formato: str) -> str:
    if isinstance(valor, (datetime, date)):
        return valor.strftime(formato)
    texto = str(valor or "").strip()
    if not texto:
        return ""
    for patron in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto[:19], patron).strftime(formato)
        except ValueError:
            continue
    return ""


def _numero_cpe(cpe: dict) -> str:
    sucursal, nro_orden = _texto(cpe, "cabecera.sucursal"), _texto(cpe, "cabecera.nroOrden")
    if not (sucursal.isdigit() and nro_orden.isdigit()):
        return ""
    return config.formato_numero_cpe.format(sucursal=int(sucursal), nro_orden=int(nro_orden))


# ---------------------------------------------------------------------------
# Fila de Visec
# ---------------------------------------------------------------------------

@dataclass
class Fila:
    ctg: str
    valores: dict[str, str]
    avisos: list[str] = field(default_factory=list)
    estado: str = ""


def armar_fila(consulta: Consulta, propios: dict[str, str] | None = None) -> Fila:
    """Traduce una CPE de ARCA a una fila de Visec.

    `propios` son los datos del sistema (los que trae el CSV de `ctg_soja.py`)
    que ARCA no tiene: la fecha del movimiento, el peso que entró a stock, los
    kilos estimados y el RENSPA.
    """
    cpe = consulta.cpe or {}
    propios = propios or {}
    avisos: list[str] = []

    fecha_movimiento = _fecha(propios.get("fecha_movimiento"), config.formato_fecha_hora)
    if not fecha_movimiento:
        fecha_movimiento = _fecha(
            _buscar(cpe, "cabecera.fechaInicioEstado") or _buscar(cpe, "cabecera.fechaEmision"),
            config.formato_fecha_hora,
        )
        avisos.append("Fecha de movimiento tomada de ARCA: verificar que sea la del ingreso a planta")

    peso_stock = _kilos_propios(propios.get("peso_ingreso_stock"))
    if not peso_stock:
        peso_stock = _kilos(cpe, "datosCarga.pesoBrutoDescarga", "datosCarga.pesoTaraDescarga")

    cod_grano = _texto(cpe, "datosCarga.codGrano")
    producto = config.grano_visec.get(cod_grano, cod_grano)

    valores = {
        "Fecha y Hora Movimiento": fecha_movimiento,
        "Fecha CPE": _fecha(_buscar(cpe, "cabecera.fechaEmision"), config.formato_fecha),
        "Número CPE": _numero_cpe(cpe),
        "Número CTG": _texto(cpe, "cabecera.nroCTG") or consulta.ctg,
        # La CPE automotor no tiene un campo "titular": es quien la emite, que
        # en la respuesta es el CUIT de origen (manual v2.0.5, respuesta de
        # autorizarCPEAutomotor).
        "CUIT Titular": _cuit(cpe, "origen.cuit"),
        "Número RUCA Origen": _texto(cpe, "origen.planta") or _texto(cpe, "origen.plantaARCA"),
        "CUIT Remitente Comercial Productor": _cuit(cpe, "retiroProductor.cuitRemitenteComercialProductor"),
        "CUIT Rte Comercial Venta Primaria": _cuit(cpe, "intervinientes.cuitRemitenteComercialVentaPrimaria"),
        "CUIT Rte Comercial Venta Secundaria": _cuit(cpe, "intervinientes.cuitRemitenteComercialVentaSecundaria"),
        "CUIT Rte Comercial Venta Secundaria 2": _cuit(cpe, "intervinientes.cuitRemitenteComercialVentaSecundaria2"),
        "CUIT Corredor Venta Primaria": _cuit(cpe, "intervinientes.cuitCorredorVentaPrimaria"),
        "CUIT Corredor Venta Secundaria": _cuit(cpe, "intervinientes.cuitCorredorVentaSecundaria"),
        "CUIT Destinatario": _cuit(cpe, "destinatario.cuit"),
        "CUIT Destino": _cuit(cpe, "destino.cuit"),
        "Número RUCA Destino": _texto(cpe, "destino.planta"),
        "Código Producto": producto,
        "Campaña": _texto(cpe, "datosCarga.cosecha"),
        "Peso Neto Carga (Kg)": _kilos(cpe, "datosCarga.pesoBruto", "datosCarga.pesoTara"),
        "Número RENSPA": (propios.get("renspa") or "").strip(),
        # Criterio de Rivara: el CTG asignado es el mismo CTG de la CPE.
        "Número CTG Asignado": _texto(cpe, "cabecera.nroCTG") or consulta.ctg,
        # Por UP (unidad productiva): la carga sale de IPL_KILOS_ESTIMADOS y la
        # descarga es el mismo peso que entró a stock.
        "Peso Neto Carga (Kg) por UP": _kilos_propios(propios.get("kilos_estimados")),
        "Peso Neto Descarga (Kg) por UP": peso_stock,
        "Peso Ingreso Stock (Kg)": peso_stock,
        "Último Almacenamiento": config.ultimo_almacenamiento,
        "Tipo Movimiento": config.tipo_movimiento,
    }

    # Lo que Visec exige y acá quedaría vacío: mejor verlo antes de subir.
    obligatorios = [
        "Fecha y Hora Movimiento", "Fecha CPE", "Número CPE", "Número CTG",
        "CUIT Titular", "CUIT Destinatario", "CUIT Destino", "Número RUCA Destino",
        "Código Producto", "Campaña", "Peso Neto Carga (Kg)", "Tipo Movimiento",
    ]
    for columna in obligatorios:
        if not valores[columna]:
            avisos.append(f"Falta {columna}")
    if not valores["Peso Ingreso Stock (Kg)"]:
        avisos.append("Falta Peso Ingreso Stock: la CPE todavía no tiene pesos de descarga en ARCA")

    codigo = consulta.estado.upper()
    estado = f"{codigo} · {ESTADOS[codigo]}" if codigo in ESTADOS else codigo
    if codigo and codigo not in ESTADOS_FINALES:
        avisos.append(f"La CPE está en estado {estado}: los pesos de descarga pueden no ser definitivos")

    return Fila(ctg=consulta.ctg, valores=valores, avisos=avisos, estado=estado)


# ---------------------------------------------------------------------------
# Escritura del XLSX
# ---------------------------------------------------------------------------

def escribir(filas: list[Fila], destino: Path) -> Path:
    libro = load_workbook(PLANTILLA)
    if HOJA not in libro.sheetnames:
        raise ValueError(f"La plantilla no tiene la hoja '{HOJA}'")
    hoja = libro[HOJA]

    encabezados = [hoja.cell(1, i + 1).value for i in range(len(COLUMNAS))]
    if encabezados != COLUMNAS:
        raise ValueError(
            "La plantilla de Visec cambió de columnas. Esperado:\n"
            f"{COLUMNAS}\nEncontrado:\n{encabezados}"
        )

    if hoja.max_row > 1:
        hoja.delete_rows(2, hoja.max_row - 1)

    # Todas las columnas de la plantilla son de texto ('@'): si se escriben como
    # número, Excel se come los ceros a la izquierda de los CUIT y los CTG.
    formatos = [hoja.cell(1, i + 1).number_format for i in range(len(COLUMNAS))]
    for numero_fila, fila in enumerate(filas, start=2):
        for i, columna in enumerate(COLUMNAS):
            celda = hoja.cell(numero_fila, i + 1, fila.valores[columna] or None)
            celda.number_format = formatos[i]

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino
