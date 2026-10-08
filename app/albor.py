"""SoftCereal -> Albor: arma la planilla de importación de comprobantes de
cosecha de Albor (ImportacionCosecha.xlsx, hoja "Cosechas a importar") a
partir de la exportación de ingresos a planta de SoftCereal.

La plantilla la baja cada uno de Albor y se sube a la app: trae la hoja
"Referencias" con las listas de códigos de Albor (campañas, especies, cultivos,
depósitos, transportistas, choferes...). No va en el repositorio: tiene nombres
y CUIT de personas.

Lo que SoftCereal llama de una forma y Albor de otra (el lote "El Bagual-lote
15" de soja ESP 25/26 es el cultivo "02836 - El Bagual LTV 15 SOJA ESP 26/27")
se resuelve con una tabla de equivalencias: la app sugiere, el usuario confirma
una vez, y queda guardada para las próximas.
"""

from __future__ import annotations

import difflib
import io
import json
import os
import re
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.sax.saxutils import escape

ALBOR_DIR = Path(os.getenv("ALBOR_DIR", "/data/albor"))
PLANTILLA = ALBOR_DIR / "plantilla_cosecha.xlsx"
REFERENCIAS = ALBOR_DIR / "referencias.json"     # las listas, leídas una vez al subir la plantilla
EQUIVALENCIAS = ALBOR_DIR / "equivalencias.json"
TANDAS = ALBOR_DIR / "tandas"                    # filas leídas, esperando equivalencias
RETENCION_TANDAS = timedelta(days=2)

HOJA_DATOS = "Cosechas a importar"
HOJA_REFERENCIAS = "Referencias"

# Cómo se escriben los valores. Todas las celdas de la plantilla son texto.
# Si Albor rechaza los códigos o los decimales, se cambia en el .env.
SOLO_CODIGO = os.getenv("ALBOR_SOLO_CODIGO", "1") != "0"   # "02836" o "02836 - El Bagual..."
SEPARADOR_DECIMAL = os.getenv("ALBOR_SEPARADOR_DECIMAL", ",")
FORMATO_FECHA = os.getenv("ALBOR_FORMATO_FECHA", "%d/%m/%Y")
TIPO_FLETE = os.getenv("ALBOR_TIPO_FLETE", "T")             # T Tercero, I Interno

# Listas de la hoja Referencias que usa la conversión: nombre en la app -> columna.
LISTAS = {
    "campana": "Código campaña",
    "especie": "Código especie",
    "grano": "Tipo de Grano",
    "cultivo": "Código cultivo",
    "deposito": "Código depósito",
    "destino": "Código destino",
    "transportista": "Código transportista",
    "chofer": "Chofer (CUIT)",
    "flete": "Tipo de Flete",
}

# Equivalencias que se confirman a mano: (tipo, título, lista de Referencias, obligatoria)
TIPOS_EQUIVALENCIA = [
    ("especie", "Especie", "especie", True),
    ("cultivo", "Cultivo (campo y lote)", "cultivo", True),
    ("deposito", "Depósito destino (por planta de SoftCereal)", "deposito", True),
    ("destino", "Destino (por planta de SoftCereal)", "destino", False),
]


class ErrorArchivo(ValueError):
    pass


def _normalizar(texto: str) -> str:
    sin_tildes = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", sin_tildes.lower()).strip()


def codigo(valor: str) -> str:
    """'02836 - El Bagual LTV 15 SOJA ESP 26/27' -> '02836'."""
    return valor.split(" - ", 1)[0].strip() if " - " in valor else valor.strip()


def _valor_lista(valor: str) -> str:
    return codigo(valor) if SOLO_CODIGO else valor.strip()


# ---------------------------------------------------------------------------
# Plantilla de Albor y sus referencias
# ---------------------------------------------------------------------------

def _rutas_hojas(libro: zipfile.ZipFile) -> dict[str, str]:
    """Nombre de hoja -> ruta del XML dentro del xlsx."""
    workbook = libro.read("xl/workbook.xml").decode("utf-8-sig")
    rels = libro.read("xl/_rels/workbook.xml.rels").decode("utf-8-sig")
    destinos = {
        m.group("id"): m.group("dst")
        for m in re.finditer(r'<(?:\w+:)?Relationship\b(?=[^>]*\bId="(?P<id>[^"]+)")(?=[^>]*\bTarget="(?P<dst>[^"]+)")', rels)
    }
    rutas = {}
    for m in re.finditer(r'<(?:\w+:)?sheet\b(?=[^>]*\bname="(?P<n>[^"]+)")(?=[^>]*\br:id="(?P<id>[^"]+)")', workbook):
        destino = destinos.get(m.group("id"), "")
        destino = destino.lstrip("/")
        rutas[_desescapar(m.group("n"))] = destino if destino.startswith("xl/") else f"xl/{destino}"
    return rutas


def _desescapar(texto: str) -> str:
    return (texto.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def es_plantilla(contenido: bytes) -> bool:
    """¿Es la plantilla de Albor? Mira solo los nombres de las hojas, así que
    sirve para reconocerla aunque la suelten en el recuadro equivocado."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            return HOJA_DATOS in _rutas_hojas(z)
    except (zipfile.BadZipFile, KeyError):
        return False


# ---------------------------------------------------------------------------
# Listado de comprobantes de Albor (lo que ya está cargado)
# ---------------------------------------------------------------------------

def _html_de_listado(contenido: bytes) -> str | None:
    """El "Excel" que exporta la grilla de comprobantes de Albor es una tabla
    HTML en UTF-16 con extensión .xls. Devuelve el HTML, o None si no es eso."""
    for codificacion in ("utf-16", "utf-8-sig"):
        if codificacion == "utf-16" and not contenido.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            texto = contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
        if "<table" in texto[:2000].lower():
            return texto
    return None


def es_listado_albor(contenido: bytes) -> bool:
    texto = _html_de_listado(contenido)
    return bool(texto) and "Nro CTG" in texto[:5000]


def ctg_en_listado(contenido: bytes) -> set[str]:
    """Los CTG de los comprobantes del listado (columnas Nro CTG y Carta Porte)."""
    import html as html_mod

    texto = _html_de_listado(contenido) or ""
    filas = [
        [html_mod.unescape(re.sub(r"<[^>]+>", "", celda)).replace("\xa0", " ").strip()
         for celda in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", fila, re.S | re.I)]
        for fila in re.findall(r"<tr[^>]*>(.*?)</tr>", texto, re.S | re.I)
    ]
    if not filas:
        return set()
    columnas = [i for i, nombre in enumerate(filas[0]) if nombre in ("Nro CTG", "Carta Porte")]
    ctgs = set()
    for fila in filas[1:]:
        for i in columnas:
            valor = re.sub(r"\D", "", fila[i]) if i < len(fila) else ""
            if len(valor) == 11:
                ctgs.add(valor)
    return ctgs


def leer_plantilla(contenido: bytes) -> dict:
    """Valida la plantilla de Albor y devuelve sus referencias."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            rutas = _rutas_hojas(z)
    except (zipfile.BadZipFile, KeyError) as e:
        raise ErrorArchivo(f"No es un .xlsx válido: {e}")
    if HOJA_DATOS not in rutas or HOJA_REFERENCIAS not in rutas:
        raise ErrorArchivo(
            f"No es la plantilla de importación de cosechas de Albor: le faltan las hojas "
            f"«{HOJA_DATOS}» y «{HOJA_REFERENCIAS}»."
        )
    libro = load_workbook(io.BytesIO(contenido), read_only=True)
    encabezado = [str(v or "").strip() for v in next(libro[HOJA_DATOS].iter_rows(max_row=1, values_only=True))]
    faltan = [c for c in COLUMNAS_USADAS if c not in encabezado]
    if faltan:
        raise ErrorArchivo("A la plantilla de Albor le faltan columnas: " + ", ".join(faltan))

    filas = libro[HOJA_REFERENCIAS].iter_rows(values_only=True)
    titulos = [str(v or "").strip() for v in next(filas)]
    indices = {clave: titulos.index(col) for clave, col in LISTAS.items() if col in titulos}
    listas: dict[str, list[str]] = {clave: [] for clave in LISTAS}
    for fila in filas:
        for clave, i in indices.items():
            if i < len(fila) and fila[i] not in (None, ""):
                listas[clave].append(str(fila[i]).strip())
    # Sin repetidos, en el orden de Albor.
    listas = {clave: list(dict.fromkeys(valores)) for clave, valores in listas.items()}
    return {"encabezado": encabezado, "listas": listas, "cargada": datetime.now().isoformat(timespec="seconds")}


def guardar_plantilla(contenido: bytes) -> dict:
    referencias = leer_plantilla(contenido)
    ALBOR_DIR.mkdir(parents=True, exist_ok=True)
    _escribir_atomico(PLANTILLA, contenido)
    _escribir_atomico(REFERENCIAS, json.dumps(referencias, ensure_ascii=False).encode())
    return referencias


def referencias() -> dict | None:
    if not (PLANTILLA.is_file() and REFERENCIAS.is_file()):
        return None
    return json.loads(REFERENCIAS.read_text(encoding="utf-8"))


def _escribir_atomico(destino: Path, contenido: bytes) -> None:
    temporal = destino.with_name(destino.name + ".tmp")
    temporal.write_bytes(contenido)
    temporal.replace(destino)


# ---------------------------------------------------------------------------
# Equivalencias SoftCereal -> Albor
# ---------------------------------------------------------------------------

def equivalencias() -> dict[str, dict[str, str]]:
    datos = json.loads(EQUIVALENCIAS.read_text(encoding="utf-8")) if EQUIVALENCIAS.is_file() else {}
    return {tipo: dict(datos.get(tipo, {})) for tipo, *_ in TIPOS_EQUIVALENCIA}


def guardar_equivalencias(nuevas: dict[str, dict[str, str]], reemplazar: bool = False) -> None:
    """Suma (o con reemplazar, pisa) equivalencias. Un valor vacío la borra."""
    actuales = {t: {} for t, *_ in TIPOS_EQUIVALENCIA} if reemplazar else equivalencias()
    for tipo, valores in nuevas.items():
        for clave, valor in valores.items():
            if valor.strip():
                actuales.setdefault(tipo, {})[clave] = valor.strip()
            else:
                actuales.setdefault(tipo, {}).pop(clave, None)
    ALBOR_DIR.mkdir(parents=True, exist_ok=True)
    _escribir_atomico(EQUIVALENCIAS, json.dumps(actuales, ensure_ascii=False, indent=1, sort_keys=True).encode())


def _sin_cosecha(descripcion: str) -> str:
    """'Soja ESP Cosecha 25/26' -> 'Soja ESP'."""
    return re.sub(r"\s*cosecha\s*\d{2}/\d{2}\s*$", "", descripcion.strip(), flags=re.I).strip()


def _palabras_especie(texto: str) -> str:
    return re.sub(r"\besp\b", "especial", _normalizar(texto))


def _especie_exacta(descripcion: str, lista: list[str]) -> str:
    buscada = _palabras_especie(_sin_cosecha(descripcion))
    for opcion in lista:
        resto = opcion.split(" - ", 1)[-1]
        if _palabras_especie(resto) == buscada:
            return opcion
    return ""


def sugerir(tipo: str, clave: str, refs: dict) -> str:
    """La opción más parecida de Albor, o '' si no hay una razonable."""
    listas = refs["listas"]
    if tipo == "especie":
        return _especie_exacta(clave, listas["especie"])
    if tipo == "cultivo":
        return _sugerir_cultivo(clave, listas["cultivo"])
    if tipo in ("deposito", "destino"):
        # Donde Albor deja hoy los ingresos de la planta de SoftCereal:
        # "PSAA - Planta Silos Acceso Alberti" (comprobante de la CPE 10134580361).
        return next((o for o in listas[tipo] if {"silo", "acceso", "alberti"} <= set(
            _normalizar(o.split(" - ", 1)[-1]).replace("silos", "silo").split())), "")
    return ""


def _numeros(texto: str) -> set[str]:
    return {n.lstrip("0") or "0" for n in re.findall(r"\d+", texto)}


def _sugerir_cultivo(clave_cultivo: str, opciones: list[str]) -> str:
    """'El Bagual-lote 15 | Soja ESP | 25/26' -> '02836 - El Bagual LTV 15 SOJA ESP 25/26'.

    Solo se sugiere un cultivo de la misma especie, cuyo nombre empiece con
    el del campo y que tenga los mismos números de lote: sugerir uno parecido
    pero equivocado es peor que no sugerir, porque se confirma sin mirar.
    Los códigos de cultivo de Albor sirven para más de una campaña (el nombre
    dice la última), así que la campaña solo desempata.
    """
    lote, especie, campana = clave_cultivo.split(" | ")
    campo, _, numero_lote = lote.partition("-")
    especie_n = _normalizar(especie).split()[0] if _normalizar(especie) else ""
    campo_n = _normalizar(campo)
    mejor, puntaje, candidatos = "", 0.0, []
    for opcion in opciones:
        descripcion = opcion.split(" - ", 1)[-1].strip()
        m = re.search(r"\s*(\d{2}/\d{2})$", descripcion)
        misma_campana = bool(m) and m.group(1) == campana
        if m:
            descripcion = descripcion[: m.start()]
        normalizada = _normalizar(descripcion)
        if especie_n not in normalizada.split() or not normalizada.startswith(campo_n):
            continue
        resto = normalizada[len(campo_n):]
        if _numeros(numero_lote) != _numeros(resto):
            continue
        # Las palabras del lote ("Entrada", "Fondo") también tienen que estar.
        palabras = set(re.findall(r"[a-z]+", _normalizar(numero_lote))) - {"lote"}
        if not palabras <= set(resto.split()):
            continue
        candidatos.append(opcion)
        parecido = difflib.SequenceMatcher(None, _normalizar(f"{lote} {especie}"), normalizada).ratio()
        parecido += 1 if misma_campana else 0
        if parecido > puntaje:
            mejor, puntaje = opcion, parecido
    # Sin lote en SoftCereal ("La Nutria-"), solo si el campo tiene un único cultivo.
    if not _normalizar(numero_lote) and len(candidatos) != 1:
        return ""
    return mejor


def _transportista(nombre: str, lista: list[str]) -> str:
    buscado = _normalizar(nombre)
    if not buscado:
        return ""
    exactos = [o for o in lista if _normalizar(o.split(" - ", 1)[-1]) == buscado]
    if exactos:
        return exactos[0]
    mejor = max(lista, key=lambda o: difflib.SequenceMatcher(None, buscado, _normalizar(o.split(" - ", 1)[-1])).ratio(), default="")
    if mejor and difflib.SequenceMatcher(None, buscado, _normalizar(mejor.split(" - ", 1)[-1])).ratio() >= 0.92:
        return mejor
    return ""


def _chofer(documento: str, lista: list[str]) -> str:
    """El DNI de SoftCereal dentro del CUIT de la lista de Albor:
    'OTTONELLO MARCOS (20-36272819-4)' -> '20-36272819-4'."""
    dni = re.sub(r"\D", "", documento)
    if not dni:
        return ""
    for opcion in lista:
        m = re.search(r"\((\d{2})-?(\d{7,8})-?(\d)\)\s*$", opcion)
        if m and m.group(2).lstrip("0") == dni.lstrip("0"):
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return ""


# ---------------------------------------------------------------------------
# Exportación de SoftCereal
# ---------------------------------------------------------------------------

# Campo interno -> encabezados posibles (normalizados) y cuál aparición tomar.
# La exportación repite "Descripción": la primera es el campo-lote, la segunda
# la especie y cosecha. "Sucursar Interna CPE" viene así, mal escrito, del ERP.
FUENTES = {
    "ctg": [("ctg", 0)],
    "ingreso": [("numero ingreso", 0), ("numero de ingreso", 0)],
    "ticket": [("numero ticket", 0), ("numero de ticket asignado", 0)],
    "planta": [("planta", 0)],
    "cosecha": [("cosecha", 0)],
    "lote": [("descripcion", 0)],
    "especie": [("descripcion", 1)],
    "fecha": [("fecha hora orden de carga", 0), ("fecha y hora calada", 0), ("fecha y hora peso bruto", 0)],
    "bruto": [("peso bruto reconocido", 0), ("peso bruto acopio", 0)],
    "tara": [("tara", 0)],
    "humedad": [("humedad", 0)],
    "estimado": [("kilos estimados", 0)],
    "km": [("kilometros acarreo", 0)],
    "transportista": [("nombre empresa de transporte", 0)],
    "chofer": [("carnet conductor transportista", 0)],
    "sucursal_cpe": [("sucursar interna cpe", 0), ("sucursal interna cpe", 0)],
    "numero_cpe": [("numero interno cpe", 0)],
    "observaciones": [("observaciones orden de carga", 0)],
}
OBLIGATORIAS = {"ctg": "CTG", "lote": "Descripción (campo-lote)", "especie": "Descripción (especie)",
                "cosecha": "Cosecha", "planta": "Planta", "bruto": "Peso Bruto Reconocido", "tara": "Tara"}


def _indices(encabezado: list[str]) -> dict[str, int]:
    apariciones: dict[str, list[int]] = {}
    for i, nombre in enumerate(encabezado):
        apariciones.setdefault(_normalizar(nombre), []).append(i)
    indices = {}
    for campo, candidatos in FUENTES.items():
        for nombre, n in candidatos:
            if len(apariciones.get(nombre, [])) > n:
                indices[campo] = apariciones[nombre][n]
                break
    faltan = [titulo for campo, titulo in OBLIGATORIAS.items() if campo not in indices]
    if faltan:
        raise ErrorArchivo(
            "El archivo no parece la exportación de ingresos de SoftCereal: falta la columna "
            + ", ".join(faltan)
        )
    return indices


def leer_softcereal(filas_planilla: list[list[str]]) -> list[dict]:
    """Filas de SoftCereal (encabezado primero) -> un dict por ingreso."""
    if len(filas_planilla) < 2:
        raise ErrorArchivo("El archivo no tiene filas de datos")
    indices = _indices(filas_planilla[0])
    ingresos = []
    # +2: los títulos de la exportación se descartan, así que se numera desde
    # el encabezado; sirve igual para ubicar el ingreso por número.
    for numero, datos in enumerate(filas_planilla[1:], start=2):
        fila = {campo: (datos[i].strip() if i < len(datos) else "") for campo, i in indices.items()}
        fila["fila"] = numero
        ingresos.append(fila)
    return ingresos


# ---------------------------------------------------------------------------
# Conversión
# ---------------------------------------------------------------------------

def _campana(cosecha: str) -> str:
    """'2526' -> '25/26'."""
    digitos = re.sub(r"\D", "", cosecha)
    return f"{digitos[:2]}/{digitos[2:]}" if len(digitos) == 4 else cosecha


def clave(tipo: str, ingreso: dict) -> str:
    if tipo == "especie":
        return _sin_cosecha(ingreso["especie"])
    if tipo == "cultivo":
        return f"{ingreso['lote']} | {_sin_cosecha(ingreso['especie'])} | {_campana(ingreso['cosecha'])}"
    return f"Planta {ingreso['planta']}"  # deposito y destino


def pendientes(ingresos: list[dict], refs: dict) -> dict[str, list[dict]]:
    """Las equivalencias que faltan confirmar, con su sugerencia.
    Las especies con nombre idéntico en Albor se guardan solas."""
    guardadas = equivalencias()
    automaticas: dict[str, dict[str, str]] = {}
    faltan: dict[str, list[dict]] = {}
    for tipo, *_ in TIPOS_EQUIVALENCIA:
        conteo: dict[str, int] = {}
        for ingreso in ingresos:
            k = clave(tipo, ingreso)
            if k and k not in guardadas[tipo]:
                conteo[k] = conteo.get(k, 0) + 1
        for k, n in sorted(conteo.items()):
            sugerencia = sugerir(tipo, k, refs)
            if tipo == "especie" and sugerencia:
                automaticas.setdefault(tipo, {})[k] = sugerencia
            else:
                faltan.setdefault(tipo, []).append({"clave": k, "sugerencia": sugerencia, "filas": n})
    if automaticas:
        guardar_equivalencias(automaticas)
    return faltan


@dataclass
class Fila:
    numero: int
    ingreso: str
    valores: dict = field(default_factory=dict)
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errores


def _numero(texto: str) -> Decimal | None:
    try:
        return Decimal((texto or "").strip().replace(",", ".")) if (texto or "").strip() else None
    except InvalidOperation:
        return None


def _texto_numero(valor: Decimal | None, decimales: int = 0) -> str:
    if valor is None:
        return ""
    if decimales == 0:
        return str(int(valor.to_integral_value()))
    texto = f"{valor:.{decimales}f}".rstrip("0").rstrip(".")
    return texto.replace(".", SEPARADOR_DECIMAL)


def _fecha(texto: str) -> str:
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto.strip(), formato).strftime(FORMATO_FECHA)
        except ValueError:
            continue
    return ""


def _armar(ingreso: dict, refs: dict, equiv: dict) -> Fila:
    listas = refs["listas"]
    fila = Fila(numero=ingreso["fila"], ingreso=ingreso.get("ingreso", ""))
    v = fila.valores

    def equivalente(tipo: str, obligatorio: bool, titulo: str) -> str:
        valor = equiv[tipo].get(clave(tipo, ingreso), "")
        if not valor and obligatorio:
            fila.errores.append(f"Falta la equivalencia de {titulo}: {clave(tipo, ingreso)}")
        return _valor_lista(valor) if valor else ""

    ctg = re.sub(r"\D", "", ingreso["ctg"])
    if len(ctg) != 11:
        fila.errores.append(f"CTG inválido: «{ingreso['ctg']}»")
    bruto, tara = _numero(ingreso["bruto"]), _numero(ingreso["tara"])
    if bruto is None or tara is None:
        fila.errores.append("Falta el peso bruto o la tara")
    elif tara >= bruto:
        fila.errores.append(f"La tara ({tara}) no es menor que el bruto ({bruto})")
    fecha = _fecha(ingreso.get("fecha", ""))
    if not fecha:
        fila.errores.append("Falta la fecha de carga")

    campana = _campana(ingreso["cosecha"])

    chofer = _chofer(ingreso.get("chofer", ""), listas["chofer"])
    transportista = _transportista(ingreso.get("transportista", ""), listas["transportista"])
    if ingreso.get("transportista") and not transportista:
        fila.avisos.append(f"Transportista sin código en Albor: {ingreso['transportista']}")
    if ingreso.get("chofer") and not chofer:
        fila.avisos.append(f"Chofer sin CUIT en Albor: DNI {ingreso['chofer']}")

    observaciones = " · ".join(p for p in (
        f"SoftCereal ingreso {ingreso['ingreso']}" if ingreso.get("ingreso") else "",
        ingreso.get("observaciones", ""),
    ) if p)[:250]

    v.update({
        "Fecha": fecha,
        "Número de ticket": ingreso.get("ticket", "").replace(" ", "-"),
        "Código campaña": campana,
        "Código especie": equivalente("especie", True, "especie"),
        # Albor tiene un solo tipo de grano ("UNICO"): si es así, va ese.
        "Tipo de Grano": _valor_lista(listas["grano"][0]) if len(listas["grano"]) == 1 else "",
        "Código cultivo": equivalente("cultivo", True, "cultivo"),
        "Código depósito destino": equivalente("deposito", True, "depósito destino"),
        "Código destino": equivalente("destino", False, "destino"),
        # Origen = destino, como lo carga hoy Albor: el productor no pesa en el campo.
        "Peso Origen Bruto": _texto_numero(bruto),
        "Peso Origen Tara": _texto_numero(tara),
        "Peso Origen Neto": _texto_numero(bruto - tara) if bruto is not None and tara is not None else "",
        "% Humedad Destino": _texto_numero(_numero(ingreso.get("humedad", "")), 2),
        "Peso Destino Bruto": _texto_numero(bruto),
        "Peso Destino Tara": _texto_numero(tara),
        "Peso Destino Neto": _texto_numero(bruto - tara) if bruto is not None and tara is not None else "",
        "Tipo de Flete": TIPO_FLETE,
        "Chofer (CUIT)": chofer,
        "Código transportista": _valor_lista(transportista) if transportista else "",
        "Tipo CPE": "E",
        # En Albor la carta de porte de una CPE es el número de CTG.
        "Carta de Porte": ctg,
        "Flete Corto": "No",
        "CTG": ctg,
        "Fecha Partida": fecha,
        "Observaciones remitente": observaciones,
        # La CPE ya existe en ARCA: que Albor no pida otra.
        "Obtener COT": "No",
        "Obtener CTG": "No",
    })
    return fila


# Columnas de Albor que llena la conversión (las demás van vacías).
COLUMNAS_USADAS = [
    "Fecha", "Número de ticket", "Código campaña", "Código especie", "Tipo de Grano",
    "Código cultivo", "Código depósito destino", "Código destino",
    "Peso Origen Bruto", "Peso Origen Tara", "Peso Origen Neto",
    "% Humedad Destino", "Peso Destino Bruto", "Peso Destino Tara", "Peso Destino Neto",
    "Tipo de Flete", "Chofer (CUIT)", "Código transportista", "Tipo CPE",
    "Carta de Porte", "Flete Corto", "CTG", "Fecha Partida",
    "Observaciones remitente", "Obtener COT", "Obtener CTG",
]


@dataclass
class Resultado:
    filas: list[Fila]
    repetidos: list[str]
    ya_en_albor: list[dict] = field(default_factory=list)

    @property
    def validas(self) -> list[Fila]:
        return [f for f in self.filas if f.ok]

    @property
    def con_error(self) -> list[Fila]:
        return [f for f in self.filas if not f.ok]

    @property
    def avisos(self) -> dict[str, int]:
        """Aviso -> cantidad de filas, para mostrarlos agrupados."""
        conteo: dict[str, int] = {}
        for f in self.validas:
            for a in f.avisos:
                conteo[a] = conteo.get(a, 0) + 1
        return dict(sorted(conteo.items(), key=lambda x: -x[1]))


def convertir(ingresos: list[dict], refs: dict, ya_en_albor: list[dict] | None = None) -> Resultado:
    equiv = equivalencias()
    filas, repetidos, vistos = [], [], set()
    for ingreso in ingresos:
        fila = _armar(ingreso, refs, equiv)
        ctg = fila.valores["CTG"]
        if fila.ok and ctg in vistos:
            repetidos.append(f"CTG {ctg} (fila {fila.numero}): repetido, va una sola vez")
            continue
        vistos.add(ctg)
        filas.append(fila)
    return Resultado(filas, repetidos, ya_en_albor or [])


def separar_cargados(ingresos: list[dict], ctgs: set[str]) -> tuple[list[dict], list[dict]]:
    """(los que faltan cargar, los que ya están en Albor)."""
    faltan, cargados = [], []
    for ingreso in ingresos:
        (cargados if re.sub(r"\D", "", ingreso["ctg"]) in ctgs else faltan).append(ingreso)
    return faltan, cargados


# ---------------------------------------------------------------------------
# Tandas: las filas leídas, mientras se confirman las equivalencias
# ---------------------------------------------------------------------------

def guardar_tanda(ingresos: list[dict], origen: str, ya_en_albor: list[dict] | None = None) -> str:
    TANDAS.mkdir(parents=True, exist_ok=True)
    limite = datetime.now() - RETENCION_TANDAS
    for vieja in TANDAS.glob("*.json"):
        if datetime.fromtimestamp(vieja.stat().st_mtime) < limite:
            vieja.unlink(missing_ok=True)
    tanda = uuid.uuid4().hex
    (TANDAS / f"{tanda}.json").write_text(json.dumps({"origen": origen, "ingresos": ingresos, "ya_en_albor": ya_en_albor or []}, ensure_ascii=False))
    return tanda


def leer_tanda(tanda: str) -> dict | None:
    if not re.fullmatch(r"[0-9a-f]{32}", tanda or ""):
        return None
    ruta = TANDAS / f"{tanda}.json"
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.is_file() else None


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

def _columna(n: int) -> str:
    letras = ""
    while n:
        n, resto = divmod(n - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def _hoja_xml(encabezado: list[str], filas: list[Fila], estilo: str) -> bytes:
    """La hoja de datos, con todas las celdas como texto: así viene la plantilla."""
    s = f' s="{estilo}"' if estilo else ""

    def celda(ref: str, texto: str, con_estilo: bool) -> str:
        atributos = s if con_estilo else ""
        if not texto:
            return f'<x:c r="{ref}"{atributos} t="inlineStr" />'
        return f'<x:c r="{ref}"{atributos} t="inlineStr"><x:is><x:t xml:space="preserve">{escape(texto)}</x:t></x:is></x:c>'

    partes = ['<?xml version="1.0" encoding="utf-8"?>'
              '<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:sheetData>']
    partes.append('<x:row r="1">' + "".join(
        celda(f"{_columna(i)}1", nombre, False) for i, nombre in enumerate(encabezado, start=1)) + "</x:row>")
    for n, fila in enumerate(filas, start=2):
        partes.append(f'<x:row r="{n}">' + "".join(
            celda(f"{_columna(i)}{n}", str(fila.valores.get(nombre, "") or ""), True)
            for i, nombre in enumerate(encabezado, start=1)) + "</x:row>")
    partes.append("</x:sheetData></x:worksheet>")
    return "".join(partes).encode("utf-8")


def escribir(filas: list[Fila], refs: dict, destino: Path) -> Path:
    """Copia la plantilla de Albor cambiando solo la hoja de datos: la hoja
    Referencias (30 MB de XML) pasa tal cual, sin abrirla."""
    with zipfile.ZipFile(PLANTILLA) as origen:
        ruta = _rutas_hojas(origen)[HOJA_DATOS]
        original = origen.read(ruta).decode("utf-8-sig")
        m = re.search(r'<(?:\w+:)?c r="[A-Z]+2"[^>]*\bs="(\d+)"', original)
        estilo = m.group(1) if m else ""
        destino.parent.mkdir(parents=True, exist_ok=True)
        temporal = destino.with_name(destino.name + ".tmp")
        with zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as salida:
            for item in origen.infolist():
                datos = _hoja_xml(refs["encabezado"], filas, estilo) if item.filename == ruta else origen.read(item)
                salida.writestr(item, datos)
        temporal.replace(destino)
    return destino
