"""Lectura de las planillas que exporta el editor SQL del ERP.

El ERP exporta a .xls (Excel 97-2003) con unos renglones de título arriba
("Exportación de datos a una Hoja de Microsoft Excel:", "Editor de SQL") y
después la fila de encabezados con los nombres de las columnas de la consulta.
La planilla se convierte a texto separado por tabulaciones desde esa fila, y de
ahí la lee el mismo código que lee los CSV: las columnas se reconocen por el
nombre, así que no importa el orden ni si se agregan columnas.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

# Firmas de los dos formatos de Excel.
OLE2 = b"\xd0\xcf\x11\xe0"  # .xls (Excel 97-2003)
ZIP = b"PK\x03\x04"         # .xlsx


def es_planilla(contenido: bytes) -> bool:
    return contenido.startswith(OLE2) or contenido.startswith(ZIP)


def _texto(valor) -> str:
    """Celda -> texto. Los números enteros sin '.0': un CTG o un peso."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        # Excel guarda la fecha como fraccion de dia: 17:33:40 puede volver como
        # 17:33:39.9999. Se redondea al segundo para no correr la hora.
        if valor.microsecond >= 500_000:
            valor += timedelta(seconds=1)
        return valor.replace(microsecond=0).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).replace("\t", " ").replace("\n", " ").strip()


def _filas_xls(contenido: bytes) -> list[list[str]]:
    import xlrd

    libro = xlrd.open_workbook(file_contents=contenido)
    hoja = libro.sheet_by_index(0)
    filas = []
    for r in range(hoja.nrows):
        fila = []
        for celda in hoja.row(r):
            if celda.ctype == xlrd.XL_CELL_DATE:
                fila.append(_texto(xlrd.xldate_as_datetime(celda.value, libro.datemode)))
            elif celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                fila.append("")  # el ERP exporta los NULL como celda de error
            else:
                fila.append(_texto(celda.value))
        filas.append(fila)
    return filas


def _filas_xlsx(contenido: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
    return [[_texto(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]


def a_texto(contenido: bytes) -> str:
    """Planilla -> texto separado por tabulaciones, desde la fila de encabezados.

    Los renglones de título que el ERP pone arriba se descartan: los datos
    arrancan en la primera fila que tiene una columna llamada CTG.
    """
    filas = _filas_xls(contenido) if contenido.startswith(OLE2) else _filas_xlsx(contenido)
    inicio = next(
        (i for i, fila in enumerate(filas) if any(c.strip().upper() == "CTG" for c in fila)),
        0,
    )
    return "\n".join("\t".join(fila) for fila in filas[inicio:] if any(c.strip() for c in fila))
