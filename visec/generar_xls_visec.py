#!/usr/bin/env python3
"""Genera el XLSX de "Registración de Cartas de Porte" de Visec a partir de
la tabla INGRESOS_PLANTA.

Dos orígenes de datos:
  --xls ARCHIVO     un extracto de INGRESOS_PLANTA exportado a Excel (para probar)
  --desde/--hasta   consulta directa a la base (DB_URL en .env)

Ejemplos:
  python generar_xls_visec.py --xls 28.xls
  python generar_xls_visec.py --desde 2026-10-01 --hasta 2026-10-01
  python generar_xls_visec.py --desde 2026-10-01 --hasta 2026-10-07 --planta 2

El resultado es una copia de plantilla_visec.xlsx con las filas cargadas, y un
archivo *_avisos.txt con los datos que faltaron y hay que revisar a mano.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
PLANTILLA = BASE / "plantilla_visec.xlsx"
QUERY = BASE / "query_ingresos.sql"


# ---------------------------------------------------------------------------
# Configuración (se lee de .env / variables de entorno)
# ---------------------------------------------------------------------------

def cargar_env(ruta: Path) -> None:
    """Carga un .env simple (CLAVE=valor) sin pisar variables ya definidas."""
    if not ruta.exists():
        return
    for linea in ruta.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#") or "=" not in linea:
            continue
        clave, valor = linea.split("=", 1)
        os.environ.setdefault(clave.strip(), valor.strip().strip('"').strip("'"))


def mapa_env(nombre: str) -> dict[str, str]:
    """Lee un mapa 'a:b,c:d' desde una variable de entorno."""
    crudo = os.environ.get(nombre, "")
    salida = {}
    for par in crudo.split(","):
        if ":" in par:
            k, v = par.split(":", 1)
            salida[k.strip()] = v.strip()
    return salida


# ---------------------------------------------------------------------------
# Lectura de datos
# ---------------------------------------------------------------------------

def leer_xls(ruta: Path) -> pd.DataFrame:
    return pd.read_excel(ruta, dtype=str)


def leer_db(desde: date, hasta: date, planta: str | None) -> pd.DataFrame:
    from sqlalchemy import create_engine, text

    url = os.environ.get("DB_URL")
    if not url:
        sys.exit("Falta DB_URL en .env (ver .env.example).")
    sql = QUERY.read_text(encoding="utf-8")
    params = {
        "desde": datetime.combine(desde, datetime.min.time()),
        # hasta inclusive: se consulta hasta el inicio del día siguiente
        "hasta": datetime.combine(hasta + timedelta(days=1), datetime.min.time()),
    }
    engine = create_engine(url)
    with engine.connect() as conn:
        df = pd.read_sql(text(sql), conn, params=params)
    df.columns = [c.upper() for c in df.columns]
    df = df.astype(str)  # v() descarta 'None'/'nan'
    if planta:
        df = df[df["IPL_PLANTA"].astype(str).str.strip() == planta]
    return df


# ---------------------------------------------------------------------------
# Helpers de formato
# ---------------------------------------------------------------------------

def v(fila: pd.Series, col: str) -> str:
    """Valor limpio de una columna ('' si no existe o es nulo)."""
    if col not in fila.index:
        return ""
    x = fila[col]
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    s = str(x).strip()
    return "" if s.lower() in ("nan", "nat", "none") else s


def cuit(s: str) -> str:
    """CUIT solo dígitos (11). Devuelve '' si no es válido o es 0."""
    d = re.sub(r"\D", "", s or "")
    return d if len(d) == 11 else ""


def entero(s: str) -> str:
    """'32840.0' -> '32840'."""
    if not s:
        return ""
    try:
        return str(int(round(float(s.replace(",", ".")))))
    except ValueError:
        return ""


def fecha_hora(s: str) -> str:
    if not s:
        return ""
    try:
        return pd.to_datetime(s).strftime(os.environ.get("FORMATO_FECHA_HORA", "%d/%m/%Y %H:%M"))
    except (ValueError, TypeError):
        return ""


def fecha(s: str) -> str:
    if not s:
        return ""
    try:
        return pd.to_datetime(s).strftime(os.environ.get("FORMATO_FECHA", "%d/%m/%Y"))
    except (ValueError, TypeError):
        return ""


# ---------------------------------------------------------------------------
# Mapeo INGRESOS_PLANTA -> columnas Visec
# ---------------------------------------------------------------------------

COLUMNAS_VISEC = [
    "Fecha y Hora Movimiento",                # A
    "Fecha CPE",                              # B
    "Número CPE",                             # C
    "Número CTG",                             # D
    "CUIT Titular",                           # E
    "Número RUCA Origen",                     # F
    "CUIT Remitente Comercial Productor",     # G
    "CUIT Rte Comercial Venta Primaria",      # H
    "CUIT Rte Comercial Venta Secundaria",    # I
    "CUIT Rte Comercial Venta Secundaria 2",  # J
    "CUIT Corredor Venta Primaria",           # K
    "CUIT Corredor Venta Secundaria",         # L
    "CUIT Destinatario",                      # M
    "CUIT Destino",                           # N
    "Número RUCA Destino",                    # O
    "Código Producto",                        # P
    "Campaña",                                # Q
    "Peso Neto Carga (Kg)",                   # R
    "Número RENSPA",                          # S
    "Número CTG Asignado",                    # T
    "Peso Neto Carga (Kg) por UP",            # U
    "Peso Neto Descarga (Kg) por UP",         # V
    "Peso Ingreso Stock (Kg)",                # W
    "Último Almacenamiento",                  # X
    "Tipo Movimiento",                        # Y
]


def mapear(fila: pd.Series, avisos: list[str]) -> dict[str, str]:
    cuit_rivara = cuit(os.environ.get("CUIT_EMPRESA", ""))
    especies = mapa_env("MAPA_ESPECIE_PRODUCTO")
    ruca_planta = mapa_env("MAPA_PLANTA_RUCA")
    tipo_mov = mapa_env("MAPA_TIPO_MOVIMIENTO")

    ctg = v(fila, "IPL_CTG") or v(fila, "IPL_NUMERO_COMP_EXTERNO")
    ref = f"CTG {ctg or '?'} (ingreso {v(fila, 'IPL_PUNTO_INGRESO')}-{v(fila, 'IPL_NUMERO_INGRESO')})"

    def falta(campo: str) -> None:
        avisos.append(f"{ref}: falta {campo}")

    # Número CPE = sucursal-número interno de la CPE
    suc, nro = entero(v(fila, "IPL_SUCURSAL_INTERNA_CPE")), entero(v(fila, "IPL_NUMERO_INTERNO_CPE"))
    numero_cpe = f"{int(suc):05d}-{int(nro):08d}" if suc and nro else ""

    # Titular: el productor; en transferencias propias (sin productor) es la empresa
    titular = cuit(v(fila, "IPL_CUIT_PRODUCTOR")) or (cuit_rivara if v(fila, "IPL_TIPO_INGRESO") == "T" else "")

    especie = v(fila, "IPL_ESPECIE")
    producto = especies.get(especie, "")
    if not producto:
        falta(f"Código Producto (especie '{especie}' sin mapear en MAPA_ESPECIE_PRODUCTO)")

    planta = entero(v(fila, "IPL_PLANTA"))
    tipo = v(fila, "IPL_TIPO_INGRESO")

    r = {
        "Fecha y Hora Movimiento": fecha_hora(v(fila, "IPL_FECHA_HORA_CONF_ARRIBO") or v(fila, "IPL_FECHA_HORA")),
        "Fecha CPE": fecha(v(fila, "IPL_FECHA_HORA_CARGA") or v(fila, "IPL_FECHA_HORA")),
        "Número CPE": numero_cpe,
        "Número CTG": ctg,
        "CUIT Titular": titular,
        "Número RUCA Origen": v(fila, "RUCA_ORIGEN"),
        "CUIT Remitente Comercial Productor": cuit(v(fila, "CUIT_REMITENTE_PRODUCTOR")),
        "CUIT Rte Comercial Venta Primaria": cuit(v(fila, "IPL_CYO1_CUIT")),
        "CUIT Rte Comercial Venta Secundaria": cuit(v(fila, "IPL_CYO2_CUIT")),
        "CUIT Rte Comercial Venta Secundaria 2": "",
        "CUIT Corredor Venta Primaria": cuit(v(fila, "CUIT_CORREDOR")),
        "CUIT Corredor Venta Secundaria": "",
        "CUIT Destinatario": cuit_rivara,
        "CUIT Destino": cuit_rivara,
        "Número RUCA Destino": ruca_planta.get(planta, ""),
        "Código Producto": producto,
        "Campaña": v(fila, "IPL_COSECHA"),
        "Peso Neto Carga (Kg)": entero(v(fila, "IPL_PESO_NETO_PRODUCTOR")),
        "Número RENSPA": v(fila, "IPL_NUMERO_RENSPA"),
        "Número CTG Asignado": "",
        "Peso Neto Carga (Kg) por UP": "",
        "Peso Neto Descarga (Kg) por UP": "",
        "Peso Ingreso Stock (Kg)": entero(v(fila, "IPL_PESO_NETO")),
        "Último Almacenamiento": "",
        "Tipo Movimiento": tipo_mov.get(tipo, ""),
    }

    if not r["Número CPE"]:
        falta("Número CPE (IPL_SUCURSAL_INTERNA_CPE / IPL_NUMERO_INTERNO_CPE vacíos)")
    if not r["CUIT Titular"]:
        falta("CUIT Titular")
    if not r["Peso Neto Carga (Kg)"]:
        falta("Peso Neto Carga (IPL_PESO_NETO_PRODUCTOR vacío)")
    if not r["Número RUCA Destino"]:
        falta(f"RUCA Destino (planta '{planta}' sin mapear en MAPA_PLANTA_RUCA)")
    if v(fila, "IPL_CORREDOR") and not r["CUIT Corredor Venta Primaria"]:
        falta(f"CUIT Corredor (código corredor {v(fila, 'IPL_CORREDOR')}, agregar JOIN en query)")
    if not r["Tipo Movimiento"]:
        falta(f"Tipo Movimiento (tipo ingreso '{tipo}' sin mapear en MAPA_TIPO_MOVIMIENTO)")
    return r


# ---------------------------------------------------------------------------
# Escritura sobre la plantilla
# ---------------------------------------------------------------------------

def escribir(filas: list[dict[str, str]], destino: Path) -> None:
    wb = load_workbook(PLANTILLA)
    ws = wb["Cartas de Porte"]

    # Verifica que la plantilla no haya cambiado de columnas
    encabezados = [ws.cell(1, i + 1).value for i in range(len(COLUMNAS_VISEC))]
    if encabezados != COLUMNAS_VISEC:
        sys.exit(f"La plantilla no tiene las columnas esperadas:\n{encabezados}")

    # Limpia filas previas (la plantilla viene vacía, pero por las dudas)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    formatos = [ws.cell(1, i + 1).number_format for i in range(len(COLUMNAS_VISEC))]
    for n, fila in enumerate(filas, start=2):
        for i, col in enumerate(COLUMNAS_VISEC):
            valor = fila[col]
            celda = ws.cell(n, i + 1, valor if valor != "" else None)
            celda.number_format = formatos[i]  # '@' = texto, igual que la plantilla
    wb.save(destino)


# ---------------------------------------------------------------------------

def main() -> None:
    cargar_env(BASE / ".env")
    if not cuit(os.environ.get("CUIT_EMPRESA", "")):
        sys.exit("Falta CUIT_EMPRESA (11 dígitos) en .env (ver .env.example).")

    ap = argparse.ArgumentParser(description="Genera el XLSX de CPE para Visec")
    ap.add_argument("--xls", type=Path, help="Extracto de INGRESOS_PLANTA (prueba sin DB)")
    ap.add_argument("--desde", type=date.fromisoformat, help="Fecha desde (AAAA-MM-DD)")
    ap.add_argument("--hasta", type=date.fromisoformat, help="Fecha hasta inclusive (AAAA-MM-DD)")
    ap.add_argument("--planta", help="Filtrar por IPL_PLANTA")
    ap.add_argument("--tipos", default=os.environ.get("TIPOS_INGRESO", ""),
                    help="Tipos de ingreso a incluir, ej: T,E,C (vacío = todos)")
    ap.add_argument("-o", "--salida", type=Path, help="Archivo de salida .xlsx")
    args = ap.parse_args()

    if args.xls:
        df = leer_xls(args.xls)
        if args.planta:
            df = df[df["IPL_PLANTA"].astype(str).str.strip() == args.planta]
        sufijo = args.xls.stem
    elif args.desde:
        hasta = args.hasta or args.desde
        df = leer_db(args.desde, hasta, args.planta)
        sufijo = f"{args.desde:%Y%m%d}_{hasta:%Y%m%d}"
    else:
        ap.error("Indicá --xls ARCHIVO o --desde FECHA")

    if args.tipos:
        tipos = {t.strip() for t in args.tipos.split(",") if t.strip()}
        df = df[df["IPL_TIPO_INGRESO"].astype(str).str.strip().isin(tipos)]

    if df.empty:
        sys.exit("No hay ingresos para los filtros indicados.")

    avisos: list[str] = []
    filas = [mapear(fila, avisos) for _, fila in df.iterrows()]

    salida = args.salida or BASE / "salida" / f"visec_cpe_{sufijo}.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    escribir(filas, salida)

    print(f"OK: {len(filas)} filas -> {salida}")
    if avisos:
        ruta_avisos = salida.with_name(salida.stem + "_avisos.txt")
        ruta_avisos.write_text("\n".join(avisos) + "\n", encoding="utf-8")
        print(f"ATENCIÓN: {len(avisos)} avisos -> {ruta_avisos}")
        for a in avisos[:10]:
            print("  -", a)
        if len(avisos) > 10:
            print(f"  ... y {len(avisos) - 10} más")


if __name__ == "__main__":
    main()
